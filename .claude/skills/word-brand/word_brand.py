#!/usr/bin/env python3
"""
TrustFlight Word Branding Tool
Rebuilds a Word document using TrustFlight brand styles from the master template.

Usage: python3 word_brand.py <data.json>
"""

import sys
import json
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = '/Users/alexcraiu/Desktop/Documents/Word templates/Basic Document.docx'

BLUE = '1E5BB5'   # Header bottom border + first data row top border
GRAY = 'D9D9D9'   # Table grid lines

TABLE_WIDTH = 10540  # twips — usable width (12240 page - 850*2 margins)


# ---------------------------------------------------------------------------
# XML helpers
# ---------------------------------------------------------------------------

def _get_tblPr(tbl_el):
    p = tbl_el.find(qn('w:tblPr'))
    if p is None:
        p = OxmlElement('w:tblPr')
        tbl_el.insert(0, p)
    return p


def fix_tbl_grid(table, widths):
    """Replace auto-generated equal tblGrid with actual column widths."""
    tbl = table._tbl
    old = tbl.find(qn('w:tblGrid'))
    if old is not None:
        tbl.remove(old)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(w))
        grid.append(col)
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is not None:
        tblPr.addnext(grid)
    else:
        tbl.insert(0, grid)


def fix_tbl_look(table):
    """Set tblLook to match template: firstRow and firstColumn conditional formats enabled."""
    tblPr = _get_tblPr(table._tbl)
    look = tblPr.find(qn('w:tblLook'))
    if look is None:
        look = OxmlElement('w:tblLook')
        tblPr.append(look)
    look.set(qn('w:val'), '04A0')
    for attr in ('w:firstRow', 'w:lastRow', 'w:firstColumn', 'w:lastColumn',
                 'w:noHBand', 'w:noVBand'):
        look.attrib.pop(qn(attr), None)
    look.set(qn('w:firstRow'), '1')
    look.set(qn('w:lastRow'), '0')
    look.set(qn('w:firstColumn'), '1')
    look.set(qn('w:lastColumn'), '0')
    look.set(qn('w:noHBand'), '0')
    look.set(qn('w:noVBand'), '1')


def apply_table_layout_fixed(table):
    """Force fixed column layout so Word respects explicit widths."""
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblLayout'))
    if old is not None:
        tblPr.remove(old)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)


def apply_table_borders(table, color=GRAY, sz=4):
    """Uniform single-line gray borders on all table edges and inside lines."""
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def apply_table_cell_margins(table, top=57, bottom=57):
    """Cell margins: only top and bottom (left/right inherit from TableNormal = 108 dxa each)."""
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblCellMar'))
    if old is not None:
        tblPr.remove(old)
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('bottom', bottom)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def apply_table_width(table, width=TABLE_WIDTH):
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblW'))
    if old is not None:
        tblPr.remove(old)
    w = OxmlElement('w:tblW')
    w.set(qn('w:w'), str(width))
    w.set(qn('w:type'), 'dxa')
    tblPr.append(w)


def apply_col_widths(table, widths):
    """Set fixed column widths (twips) across every row."""
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i >= len(widths):
                break
            tc = cell._tc
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = OxmlElement('w:tcPr')
                tc.insert(0, tcPr)
            old = tcPr.find(qn('w:tcW'))
            if old is not None:
                tcPr.remove(old)
            w = OxmlElement('w:tcW')
            w.set(qn('w:w'), str(widths[i]))
            w.set(qn('w:type'), 'dxa')
            tcPr.append(w)


def style_header_cell(cell, color=BLUE, sz=12):
    """
    Header row cells: nil borders on top/left/right (creates borderless header band),
    thick blue bottom border, vAlign=bottom.
    """
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)

    old = tcPr.find(qn('w:tcBorders'))
    if old is not None:
        tcPr.remove(old)
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        tcBorders.append(el)
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), str(sz))
    bot.set(qn('w:space'), '0')
    bot.set(qn('w:color'), color)
    tcBorders.append(bot)
    tcPr.append(tcBorders)

    old_va = tcPr.find(qn('w:vAlign'))
    if old_va is not None:
        tcPr.remove(old_va)
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'bottom')
    tcPr.append(va)


def set_first_data_row_border(cell, color=BLUE, sz=12):
    """First data row: thick blue top border on every cell (visual separator under header)."""
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    old = tcBorders.find(qn('w:top'))
    if old is not None:
        tcBorders.remove(old)
    el = OxmlElement('w:top')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), '0')
    el.set(qn('w:color'), color)
    tcBorders.append(el)


# ---------------------------------------------------------------------------
# Document structure
# ---------------------------------------------------------------------------

def clear_body(doc):
    """Remove all body content while preserving sectPr (page layout + headers/footers)."""
    body = doc.element.body
    sectPr = body.find(qn('w:sectPr'))
    for child in list(body):
        body.remove(child)
    if sectPr is not None:
        body.append(sectPr)


def default_col_widths(n):
    if n == 1:
        return [TABLE_WIDTH]
    if n == 2:
        return [5270, 5270]
    if n == 3:
        return [4200, 3170, 3170]  # field label 40%, two data cols 30% each
    per = TABLE_WIDTH // n
    return [per] * n


def add_section(doc, section):
    heading = section.get('heading', '')
    level = section.get('heading_level', 2)

    if heading:
        doc.add_paragraph(heading, style=f'heading {level}')

    if 'table' in section:
        t = section['table']
        headers = t['headers']
        rows = t['rows']
        widths = t.get('col_widths') or default_col_widths(len(headers))

        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = 'Table Grid'

        fix_tbl_grid(table, widths)
        fix_tbl_look(table)
        apply_table_layout_fixed(table)
        apply_table_borders(table)
        apply_table_cell_margins(table)
        apply_table_width(table, sum(widths))
        apply_col_widths(table, widths)

        # Header row: TableHeading style, nil side borders, thick blue bottom, vAlign bottom
        for i, txt in enumerate(headers):
            cell = table.rows[0].cells[i]
            para = cell.paragraphs[0]
            para.clear()
            para.style = 'Table Heading'
            para.add_run(txt)
            style_header_cell(cell)

        # Data rows: Normal style; first column bold (inline run); first row gets blue top border
        for r, row_data in enumerate(rows):
            for c, txt in enumerate(row_data):
                cell = table.rows[r + 1].cells[c]
                para = cell.paragraphs[0]
                para.clear()
                run = para.add_run(txt)
                if c == 0:
                    run.bold = True
                if r == 0:
                    set_first_data_row_border(cell)

        doc.add_paragraph()

    elif 'paragraphs' in section:
        for text in section['paragraphs']:
            doc.add_paragraph(text)


def build(data):
    doc = Document(TEMPLATE)
    clear_body(doc)

    if title := data.get('title'):
        doc.add_paragraph(title, style='heading 1')

    if subtitle := data.get('subtitle'):
        doc.add_paragraph(subtitle, style='heading 3')

    for section in data.get('sections', []):
        add_section(doc, section)

    out = data['output_path']
    doc.save(out)
    print(f'Saved: {out}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python3 word_brand.py <data.json>')
        sys.exit(1)
    with open(sys.argv[1]) as f:
        build(json.load(f))
