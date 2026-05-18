#!/usr/bin/env python3
"""
TrustFlight Letterhead Tool
Generates a 1-2 page branded letter from the office-specific master templates
bundled under skills/letterhead/templates/.

Usage: python3 letterhead.py <data.json>
"""

import json
import os
import sys
from datetime import date
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

SAPPHIRE = RGBColor(0x1E, 0x5B, 0xB5)
GRAPHITE = RGBColor(0x24, 0x2D, 0x41)

SKILL_DIR = Path(__file__).resolve().parent
TEMPLATES_DIR = SKILL_DIR / 'templates'

SUPPORTED_BRANDS = ['TrustFlight', 'Baines Simmons', 'Kenyon', 'Redline']
SUPPORTED_OFFICES = ['Bracknell', 'Doncaster', 'Houston', 'Vancouver', 'London', 'Luton', 'Jersey']
BRANDS_WITH_TEMPLATES = ['TrustFlight']


def resolve_template(brand: str, office: str) -> Path:
    if brand not in SUPPORTED_BRANDS:
        raise SystemExit(f"Unknown brand '{brand}'. Allowed: {', '.join(SUPPORTED_BRANDS)}")
    if office not in SUPPORTED_OFFICES:
        raise SystemExit(f"Unknown office '{office}'. Allowed: {', '.join(SUPPORTED_OFFICES)}")
    if brand not in BRANDS_WITH_TEMPLATES:
        raise SystemExit(
            f"No letterhead template exists for {brand} yet. "
            f"Brands with templates: {', '.join(BRANDS_WITH_TEMPLATES)}."
        )
    for ext in ('.docx', '.dotx'):
        path = TEMPLATES_DIR / f'Letterhead [{office}]{ext}'
        if path.exists():
            return path
    raise SystemExit(
        f"Template not found: {TEMPLATES_DIR / f'Letterhead [{office}].(docx|dotx)'}\n"
        f"See {TEMPLATES_DIR / 'README.md'} for one-time setup."
    )


def clear_body(doc):
    """Remove all body content, preserve sectPr (page layout + headers/footers)."""
    body = doc.element.body
    sectPr = body.find(qn('w:sectPr'))
    for child in list(body):
        body.remove(child)
    if sectPr is not None:
        body.append(sectPr)


def add_paragraph(doc, text, *, bold=False, size_pt=8.5, color=None, tight=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    if tight:
        p.paragraph_format.line_spacing = 1.0
    run = p.add_run(text)
    run.font.name = 'Open Sans'
    run.font.size = Pt(size_pt)
    if bold:
        run.bold = True
    if color is not None:
        run.font.color.rgb = color
    return p


def add_blank(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)


def format_today() -> str:
    today = date.today()
    return f'{today.day} {today.strftime("%B")} {today.year}'


def build(data):
    brand = data.get('brand', 'TrustFlight')
    office = data['office']
    template = resolve_template(brand, office)

    doc = Document(str(template))
    clear_body(doc)

    add_paragraph(doc, data.get('date') or format_today())
    add_blank(doc)

    recipient = data.get('recipient') or {}
    if name := recipient.get('name'):
        add_paragraph(doc, name, bold=True, tight=True)

    role_parts = [recipient.get('title'), recipient.get('company')]
    role_line = ', '.join(part.strip() for part in role_parts if part and part.strip())
    if role_line:
        add_paragraph(doc, role_line, tight=True)

    address = recipient.get('address')
    if isinstance(address, list):
        address_line = ', '.join(part.strip() for part in address if part and part.strip())
    elif isinstance(address, str):
        address_line = address.strip()
    else:
        address_line = ''
    if address_line:
        add_paragraph(doc, address_line, tight=True)

    if email := recipient.get('email'):
        add_paragraph(doc, email, tight=True)
    if phone := recipient.get('phone'):
        add_paragraph(doc, phone, tight=True)

    if recipient:
        add_blank(doc)

    if subject := data.get('subject'):
        add_paragraph(doc, subject, bold=True)
        add_blank(doc)

    add_paragraph(doc, data.get('salutation') or 'Dear Sir/Madam,')
    add_blank(doc)

    for para in data.get('body') or []:
        add_paragraph(doc, para)
        add_blank(doc)

    add_blank(doc)
    add_paragraph(doc, data.get('signoff') or 'Kind regards,', color=GRAPHITE, tight=True)

    signer = data.get('signer') or {}
    if name := signer.get('name'):
        add_paragraph(doc, name, bold=True, color=SAPPHIRE, tight=True)
    if title := signer.get('title'):
        add_paragraph(doc, title, color=GRAPHITE, tight=True)

    out = data['output_path']
    os.makedirs(os.path.dirname(out), exist_ok=True)
    doc.save(out)
    print(f'Saved: {out}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python3 letterhead.py <data.json>')
        sys.exit(1)
    with open(sys.argv[1]) as f:
        build(json.load(f))
