#!/usr/bin/env python3
"""
TrustFlight Word Branding Tool
Rebuilds a Word document using TrustFlight brand styles from the master template.

Usage: python3 word_docs.py <data.json>
"""

import os
import sys
import json
from docx import Document
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

TEMPLATE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    'templates', 'Basic Document.docx',
)
PROPOSAL_TEMPLATE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    'templates', 'Proposal Template.docx',
)

PRODUCT_SECTIONS = {
    'tech_log':    'Tech Log Technical Overview',
    'centrik_5':   'Centrik 5 Technical Overview',
    'smart_suite': 'Smart Suite Technical Overview',
}

BLUE = '1E5BB5'   # Header bottom border + first data row top border
GRAY = 'D9D9D9'   # Table grid lines

TABLE_WIDTH = 10540  # twips — usable width (12240 page - 850*2 margins)


# ---------------------------------------------------------------------------
# XML helpers
# ---------------------------------------------------------------------------

def _get_tblPr(tbl_el):
    p = tbl_el.find(qn('w:tblPr'))
    if p is None:
        p = OxmlElement('w:tblPr')
        tbl_el.insert(0, p)
    return p


def fix_tbl_grid(table, widths):
    """Replace auto-generated equal tblGrid with actual column widths."""
    tbl = table._tbl
    old = tbl.find(qn('w:tblGrid'))
    if old is not None:
        tbl.remove(old)
    grid = OxmlElement('w:tblGrid')
    for w in widths:
        col = OxmlElement('w:gridCol')
        col.set(qn('w:w'), str(w))
        grid.append(col)
    tblPr = tbl.find(qn('w:tblPr'))
    if tblPr is not None:
        tblPr.addnext(grid)
    else:
        tbl.insert(0, grid)


def fix_tbl_look(table):
    """Set tblLook to match template: firstRow and firstColumn conditional formats enabled."""
    tblPr = _get_tblPr(table._tbl)
    look = tblPr.find(qn('w:tblLook'))
    if look is None:
        look = OxmlElement('w:tblLook')
        tblPr.append(look)
    look.set(qn('w:val'), '04A0')
    for attr in ('w:firstRow', 'w:lastRow', 'w:firstColumn', 'w:lastColumn',
                 'w:noHBand', 'w:noVBand'):
        look.attrib.pop(qn(attr), None)
    look.set(qn('w:firstRow'), '1')
    look.set(qn('w:lastRow'), '0')
    look.set(qn('w:firstColumn'), '1')
    look.set(qn('w:lastColumn'), '0')
    look.set(qn('w:noHBand'), '0')
    look.set(qn('w:noVBand'), '1')


def apply_table_layout_fixed(table):
    """Force fixed column layout so Word respects explicit widths."""
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblLayout'))
    if old is not None:
        tblPr.remove(old)
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    tblPr.append(layout)


def apply_table_borders(table, color=GRAY, sz=4):
    """Uniform single-line gray borders on all table edges and inside lines."""
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblBorders'))
    if old is not None:
        tblPr.remove(old)
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), str(sz))
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    tblPr.append(borders)


def apply_table_cell_margins(table, top=57, bottom=57):
    """Cell margins: only top and bottom (left/right inherit from TableNormal = 108 dxa each)."""
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblCellMar'))
    if old is not None:
        tblPr.remove(old)
    mar = OxmlElement('w:tblCellMar')
    for side, val in (('top', top), ('bottom', bottom)):
        el = OxmlElement(f'w:{side}')
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)
    tblPr.append(mar)


def apply_table_width(table, width=TABLE_WIDTH):
    tblPr = _get_tblPr(table._tbl)
    old = tblPr.find(qn('w:tblW'))
    if old is not None:
        tblPr.remove(old)
    w = OxmlElement('w:tblW')
    w.set(qn('w:w'), str(width))
    w.set(qn('w:type'), 'dxa')
    tblPr.append(w)


def apply_col_widths(table, widths):
    """Set fixed column widths (twips) across every row."""
    for row in table.rows:
        for i, cell in enumerate(row.cells):
            if i >= len(widths):
                break
            tc = cell._tc
            tcPr = tc.find(qn('w:tcPr'))
            if tcPr is None:
                tcPr = OxmlElement('w:tcPr')
                tc.insert(0, tcPr)
            old = tcPr.find(qn('w:tcW'))
            if old is not None:
                tcPr.remove(old)
            w = OxmlElement('w:tcW')
            w.set(qn('w:w'), str(widths[i]))
            w.set(qn('w:type'), 'dxa')
            tcPr.append(w)


def style_header_cell(cell, color=BLUE, sz=12):
    """
    Header row cells: nil borders on top/left/right (creates borderless header band),
    thick blue bottom border, vAlign=bottom.
    """
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)

    old = tcPr.find(qn('w:tcBorders'))
    if old is not None:
        tcPr.remove(old)
    tcBorders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'right'):
        el = OxmlElement(f'w:{edge}')
        el.set(qn('w:val'), 'nil')
        tcBorders.append(el)
    bot = OxmlElement('w:bottom')
    bot.set(qn('w:val'), 'single')
    bot.set(qn('w:sz'), str(sz))
    bot.set(qn('w:space'), '0')
    bot.set(qn('w:color'), color)
    tcBorders.append(bot)
    tcPr.append(tcBorders)

    old_va = tcPr.find(qn('w:vAlign'))
    if old_va is not None:
        tcPr.remove(old_va)
    va = OxmlElement('w:vAlign')
    va.set(qn('w:val'), 'bottom')
    tcPr.append(va)


def set_first_data_row_border(cell, color=BLUE, sz=12):
    """First data row: thick blue top border on every cell (visual separator under header)."""
    tc = cell._tc
    tcPr = tc.find(qn('w:tcPr'))
    if tcPr is None:
        tcPr = OxmlElement('w:tcPr')
        tc.insert(0, tcPr)
    tcBorders = tcPr.find(qn('w:tcBorders'))
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    old = tcBorders.find(qn('w:top'))
    if old is not None:
        tcBorders.remove(old)
    el = OxmlElement('w:top')
    el.set(qn('w:val'), 'single')
    el.set(qn('w:sz'), str(sz))
    el.set(qn('w:space'), '0')
    el.set(qn('w:color'), color)
    tcBorders.append(el)


# ---------------------------------------------------------------------------
# Document structure
# ---------------------------------------------------------------------------

def clear_body(doc):
    """Remove all body content while preserving sectPr (page layout + headers/footers)."""
    body = doc.element.body
    sectPr = body.find(qn('w:sectPr'))
    for child in list(body):
        body.remove(child)
    if sectPr is not None:
        body.append(sectPr)


def default_col_widths(n):
    if n == 1:
        return [TABLE_WIDTH]
    if n == 2:
        return [5270, 5270]
    if n == 3:
        return [4200, 3170, 3170]  # field label 40%, two data cols 30% each
    per = TABLE_WIDTH // n
    return [per] * n


def add_section(doc, section):
    heading = section.get('heading', '')
    level = section.get('heading_level', 2)

    if heading:
        doc.add_paragraph(heading, style=f'heading {level}')

    if 'table' in section:
        t = section['table']
        headers = t['headers']
        rows = t['rows']
        widths = t.get('col_widths') or default_col_widths(len(headers))

        table = doc.add_table(rows=1 + len(rows), cols=len(headers))
        table.style = 'Table Grid'

        fix_tbl_grid(table, widths)
        fix_tbl_look(table)
        apply_table_layout_fixed(table)
        apply_table_borders(table)
        apply_table_cell_margins(table)
        apply_table_width(table, sum(widths))
        apply_col_widths(table, widths)

        # Header row: TableHeading style, nil side borders, thick blue bottom, vAlign bottom
        for i, txt in enumerate(headers):
            cell = table.rows[0].cells[i]
            para = cell.paragraphs[0]
            para.clear()
            para.style = 'Table Heading'
            para.add_run(txt)
            style_header_cell(cell)

        # Data rows: Normal style; first column bold (inline run); first row gets blue top border
        for r, row_data in enumerate(rows):
            for c, txt in enumerate(row_data):
                cell = table.rows[r + 1].cells[c]
                para = cell.paragraphs[0]
                para.clear()
                run = para.add_run(txt)
                if c == 0:
                    run.bold = True
                if r == 0:
                    set_first_data_row_border(cell)

        doc.add_paragraph()

    elif 'paragraphs' in section:
        for text in section['paragraphs']:
            doc.add_paragraph(text)


def build(data):
    doc = Document(TEMPLATE)
    body = doc.element.body
    is_brief = bool(data.get('is_brief', False))

    rev_start, rev_end = _find_revision_history_bounds(body)
    end_page_start = _find_end_page_start(body)

    _substitute_cover_placeholders(body, data)

    detached_end_page = _detach_end_page(body, end_page_start)
    _strip_between_rev_history_and_end(body, rev_end)
    if is_brief:
        _remove_revision_history(body, rev_start, rev_end)

    for section in data.get('sections', []):
        add_section(doc, section)

    _reattach_end_page(body, detached_end_page)

    out = data['output_path']
    doc.save(out)
    print(f'Saved: {out}')


# ---------------------------------------------------------------------------
# Basic flow helpers — preserve template cover, revision history, end page
# ---------------------------------------------------------------------------

def _paragraph_text_of(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def _has_page_break(el):
    if el.tag != qn('w:p'):
        return False
    for br in el.iter(qn('w:br')):
        if br.get(qn('w:type')) == 'page':
            return True
    return False


def _find_revision_history_bounds(body):
    """Return (start_idx, end_idx_inclusive) of the Revision History block, or (None, None).
    The block runs from the 'Revision History' heading through the next page break paragraph
    (which terminates the page in the template)."""
    children = list(body)
    start = None
    for i, el in enumerate(children):
        if el.tag == qn('w:p') and _paragraph_text_of(el).strip() == 'Revision History':
            start = i
            break
    if start is None:
        return None, None
    for j in range(start + 1, len(children)):
        if children[j].tag == qn('w:sectPr'):
            return start, j - 1
        if _has_page_break(children[j]):
            return start, j
    return start, len(children) - 1


def _find_end_page_start(body):
    """The end/back page block begins at the last table in the body (the contact card)."""
    children = list(body)
    for i in range(len(children) - 1, -1, -1):
        if children[i].tag == qn('w:tbl'):
            return i
    return None


def _substitute_cover_placeholders(body, data):
    mapping = {}
    if title := data.get('title'):
        mapping['Document Title Goes Here'] = title
    if subtitle := data.get('subtitle'):
        mapping['Document Subtitle'] = subtitle
    if not mapping:
        return
    for t in body.iter(qn('w:t')):
        if t.text:
            new = t.text
            for k, v in mapping.items():
                if k in new:
                    new = new.replace(k, v)
            if new != t.text:
                t.text = new


def _detach_end_page(body, end_page_start):
    """Remove and return the end page elements (everything from end_page_start up to but
    not including sectPr). They get reattached after user content is added."""
    if end_page_start is None:
        return []
    children = list(body)
    detached = []
    for el in children[end_page_start:]:
        if el.tag == qn('w:sectPr'):
            break
        detached.append(el)
        body.remove(el)
    return detached


def _strip_between_rev_history_and_end(body, rev_end):
    """Remove the template's sample body content that sits between the Revision History
    page and the end page (which has already been detached)."""
    children = list(body)
    start = (rev_end + 1) if rev_end is not None else None
    if start is None:
        return
    end = len(children)
    for i, el in enumerate(children):
        if el.tag == qn('w:sectPr'):
            end = i
            break
    for el in children[start:end]:
        body.remove(el)


def _remove_revision_history(body, rev_start, rev_end):
    if rev_start is None or rev_end is None:
        return
    children = list(body)
    for el in children[rev_start:rev_end + 1]:
        body.remove(el)


def _reattach_end_page(body, detached):
    if not detached:
        return
    sectPr = body.find(qn('w:sectPr'))
    for el in detached:
        if sectPr is not None:
            sectPr.addprevious(el)
        else:
            body.append(el)


# ---------------------------------------------------------------------------
# Proposal flow
# ---------------------------------------------------------------------------

def _is_heading1(el):
    if el.tag != qn('w:p'):
        return False
    pPr = el.find(qn('w:pPr'))
    if pPr is None:
        return False
    pStyle = pPr.find(qn('w:pStyle'))
    if pStyle is None:
        return False
    return pStyle.get(qn('w:val')) in ('Heading1', 'heading1', 'heading 1', 'Heading 1')


def _paragraph_text(el):
    return ''.join(t.text or '' for t in el.iter(qn('w:t')))


def _remove_section_by_heading(body, heading_text):
    """Remove an H1 paragraph and every element after it up to (but not including) the next H1."""
    children = list(body)
    start = None
    for i, el in enumerate(children):
        if _is_heading1(el) and _paragraph_text(el).strip() == heading_text:
            start = i
            break
    if start is None:
        return False

    end = len(children)
    for j in range(start + 1, len(children)):
        if _is_heading1(children[j]) and _paragraph_text(children[j]).strip():
            end = j
            break

    for el in children[start:end]:
        body.remove(el)
    return True


def _replace_placeholders(doc, replacements):
    """Substring-replace inside every w:t run across body, headers, and footers."""
    if not replacements:
        return

    def walk(root_el):
        for t in root_el.iter(qn('w:t')):
            if t.text:
                new = t.text
                for needle, value in replacements.items():
                    if needle in new:
                        new = new.replace(needle, value)
                if new != t.text:
                    t.text = new

    walk(doc.element)
    for section in doc.sections:
        for hdr in (section.header, section.first_page_header, section.even_page_header):
            walk(hdr._element)
        for ftr in (section.footer, section.first_page_footer, section.even_page_footer):
            walk(ftr._element)


def build_proposal(data):
    products = data.get('products') or []
    if not products:
        print('Error: proposal requires at least one product in "products".')
        sys.exit(1)
    unknown = [p for p in products if p not in PRODUCT_SECTIONS]
    if unknown:
        print(f'Error: unknown product(s): {unknown}. Allowed: {list(PRODUCT_SECTIONS)}')
        sys.exit(1)

    doc = Document(PROPOSAL_TEMPLATE)

    excluded = [key for key in PRODUCT_SECTIONS if key not in products]
    for key in excluded:
        heading = PRODUCT_SECTIONS[key]
        removed = _remove_section_by_heading(doc.element.body, heading)
        if not removed:
            print(f'Warning: could not locate section "{heading}" to remove.')

    _replace_placeholders(doc, data.get('placeholders', {}))

    out = data['output_path']
    doc.save(out)
    print(f'Saved: {out}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        print('Usage: python3 word_docs.py <data.json>')
        sys.exit(1)
    with open(sys.argv[1]) as f:
        data = json.load(f)
    if data.get('doc_type') == 'proposal':
        build_proposal(data)
    else:
        build(data)
